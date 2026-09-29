import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
from PIL import Image, ImageDraw

def ensure_logo():
    os.makedirs("assets", exist_ok=True)
    path = "assets/logo.png"
    if not os.path.exists(path):
        img = Image.new('RGB', (400, 100), (30,30,47))
        d = ImageDraw.Draw(img)
        d.text((20,35), "LegalEase - Legal AI", fill=(255,255,255))
        img.save(path)
    return path

def sanitize_text(text):
    if not text: return ""
    for k,v in {"**":"", "###":"", "##":"", "’":"'", "“":'"', "”":'"'}.items():
        text = text.replace(k,v)
    return text.strip()

def format_docx(text, doc_type, parties="", terms_raw=""):
    ensure_logo()
    doc = Document()
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(ensure_logo(), width=Inches(1.5))
    except: pass

    title = doc.add_heading(doc_type, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(sanitize_text(text))

    if terms_raw:
        doc.add_heading("Terms Table", 1)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Light Grid Accent 1'
        table.rows[0].cells[0].text = "No"
        table.rows[0].cells[1].text = "Term"
        for i,t in enumerate([x for x in terms_raw.split(';') if x.strip()],1):
            r = table.add_row().cells
            r[0].text = str(i)
            r[1].text = t.strip()

    doc.add_paragraph("\nLegalEase Inc. | All Rights Reserved", style='Intense Quote')
    name = f"{doc_type.replace(' ','_')}.docx"
    doc.save(name)
    return name

def format_pdf(text, doc_type):
    ensure_logo()
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial","B",16)
    pdf.cell(0,10,"LegalEase", new_x="LMARGIN", new_y="NEXT", align='C')
    pdf.set_font("Arial","B",12)
    pdf.cell(0,10,doc_type, new_x="LMARGIN", new_y="NEXT", align='C')
    pdf.ln(5)
    pdf.set_font("Arial","",11)
    pdf.multi_cell(0,7,sanitize_text(text))
    pdf.set_y(-15)
    pdf.set_font("Arial","I",8)
    pdf.cell(0,10,"LegalEase Inc. | All Rights Reserved", align='C')
    name = f"{doc_type.replace(' ','_')}.pdf"
    pdf.output(name)
    return name